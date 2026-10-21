from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from PIL import Image, ImageEnhance
from pathlib import Path
out=Path(r'D:/HMS/Hospital-Management-System/output/cover-page')
im=Image.open(r'C:/Users/laksh/AppData/Local/Temp/codex-clipboard-deae341e-2eb7-4920-bed7-0b9d6d3f45a3.png').convert('RGB')
im.crop((76,70,460,139)).save(out/'banner.png')
crest=im.crop((77,70,144,139)).resize((700,720))
Image.blend(Image.new('RGB',crest.size,'white'),crest,.065).save(out/'watermark.png')
doc=Document(); s=doc.sections[0]
s.page_width=Inches(8.5);s.page_height=Inches(11.075)
s.top_margin=Inches(.52);s.bottom_margin=Inches(.5);s.left_margin=Inches(1.37);s.right_margin=Inches(.83)
normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(12);normal.paragraph_format.space_after=Pt(0)
borders=OxmlElement('w:pgBorders');borders.set(qn('w:offsetFrom'),'page')
for edge in ['top','left','bottom','right']:
 e=OxmlElement('w:'+edge)
 for k,v in {'val':'single','sz':'6','space':'24','color':'000000'}.items():e.set(qn('w:'+k),v)
 borders.append(e)
s._sectPr.append(borders)
# Floating image behind the editable document body.
h=s.header.paragraphs[0];r=h.add_run();r.add_picture(str(out/'watermark.png'),width=Inches(6.0),height=Inches(6.1))
inl=r._r.xpath('.//wp:inline')[0]
a=OxmlElement('wp:anchor')
for k,v in {'distT':'0','distB':'0','distL':'0','distR':'0','simplePos':'0','relativeHeight':'0','behindDoc':'1','locked':'0','layoutInCell':'1','allowOverlap':'1'}.items():a.set(k,v)
sp=OxmlElement('wp:simplePos');sp.set('x','0');sp.set('y','0');a.append(sp)
for dim,pos in [('H',1.25),('V',2.4)]:
 p=OxmlElement('wp:position'+dim);p.set('relativeFrom','page');po=OxmlElement('wp:posOffset');po.text=str(int(pos*914400));p.append(po);a.append(p)
for child in list(inl):a.append(child)
a.insert(4,OxmlElement('wp:wrapNone'));inl.getparent().replace(inl,a)
# Coordinates follow the supplied page, scaled to this page size.
scale=612/505; current=31

def line(text,y,size=12,bold=True,underline=False,color=None):
 global current
 p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 f=p.paragraph_format;f.space_before=Pt(max(0,y*scale-current));f.line_spacing=Pt(size*1.18)
 r=p.add_run(text);r.bold=bold;r.underline=underline;r.font.size=Pt(size)
 if color:r.font.color.rgb=RGBColor.from_string(color)
 current=y*scale+size*1.18
 return p
line('HOSPITAL MANAGEMENT SYSTEM',35,6.5,False,color='B7B7B7').alignment=WD_ALIGN_PARAGRAPH.RIGHT
p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_before=Pt(30)
p.add_run().add_picture(str(out/'banner.png'),width=Inches(6.46),height=Inches(1.16));current=139*scale
line('LJ INSTITUTE OF COMPUTER APPLICATION',186,11.5)
line('&',216,11.5)
line('Master of Computer Application Sem 3',245,11.5)
line('SOFTWARE PROJECT :-1',280,10,True,True)
line('PROJECT REPORT ON',320,9.5)
line('HOSPITAL MANAGEMENT SYSTEM',347,15,True)
line('Under the Guidance of',396,10.5,True,True)
line('Prof. AJAY CHAUHAN',423,10)
line('GROUP NUMBER:-63',472,10,True,True)
line('SUBMITTED BY:-',506,10,True,True)
p=doc.add_paragraph();p.paragraph_format.space_before=Pt(16);p.paragraph_format.line_spacing=Pt(1);p.paragraph_format.space_after=Pt(0);p.add_run().font.size=Pt(1)
t=doc.add_table(rows=4,cols=3);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
widths=[1.01,2.16,3.1]
rows=[['Div\nRoll No.','Enrollment\nNo.','Name'],['D45','25004401110226','Lakshit Suthar'],['D23','25004401110192','Nirmeet Makwana'],['D48','25004401110211','Princi Trapas iya']]
rows[3][2]='Princi Trapasiy a'.replace(' ','')
for i,row in enumerate(t.rows):
 for j,c in enumerate(row.cells):
  c.width=Inches(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;c.text=rows[i][j]
  for p in c.paragraphs:
   p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.line_spacing=1;p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2)
   for r in p.runs:r.font.name='Arial';r.font.size=Pt(9 if i==0 else 7.5);r.bold=i==0
  pr=c._tc.get_or_add_tcPr();bs=OxmlElement('w:tcBorders')
  for edge in ['top','left','bottom','right']:
   e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'6');e.set(qn('w:color'),'000000');bs.append(e)
  pr.append(bs)
f=s.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER;rr=f.add_run('1');rr.font.size=Pt(7)
s.footer_distance=Inches(.6)
doc.core_properties.title='Hospital Management System Project Cover Page'
doc.save(out/'Hospital Management System Cover Page.docx')
print(out/'Hospital Management System Cover Page.docx')

