from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
from PIL import Image,ImageDraw,ImageFont
import textwrap

root=Path(__file__).resolve().parent
d=Document(root.parent/'MCA_Project_Report.docx')
im=Image.new('RGB',(1800,1300),'white');draw=ImageDraw.Draw(im)
font=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',25)
bold=ImageFont.truetype('C:/Windows/Fonts/arialbd.ttf',28)
draw.rectangle((190,60,1610,1260),outline='black',width=3)
draw.text((900,85),'Hospital Management System',font=bold,anchor='mm',fill='black')
groups=[('Patient',100,360,540,160,['Register Account','Browse Doctors and Departments','Book Appointment','View and Cancel Own Appointments','View Bills and Make Demo Payment','Download Own Receipts']),('Receptionist',100,960,540,760,['Register Patient','View Patient Details','Book Patient Appointment','Manage Appointment Status','View Doctor Availability','Record Payment and Issue Receipt']),('Doctor',1700,360,1260,160,['View Appointments','View Patient Details','Create and Review Prescriptions','Manage Availability','Complete or Cancel Appointment']),('Administrator',1700,960,1260,720,['Manage Doctors','Manage Receptionists and Shifts','View Patient Records','Monitor Appointments','Manage Billing','Manage User Access and Security','Generate Reports'])]
for role,ax,ay,cx,start,labels in groups:
    draw.ellipse((ax-18,ay-75,ax+18,ay-39),outline='black',width=3)
    draw.line((ax,ay-39,ax,ay+25),fill='black',width=3)
    draw.line((ax-34,ay-12,ax+34,ay-12),fill='black',width=3)
    draw.line((ax,ay+25,ax-30,ay+66),fill='black',width=3)
    draw.line((ax,ay+25,ax+30,ay+66),fill='black',width=3)
    draw.text((ax,ay+95),role,font=font,anchor='mm',fill='black')
    for j,label in enumerate(labels):
        y=start+j*78
        edge=cx-265 if ax<cx else cx+265
        draw.line((ax,ay,edge,y),fill='black',width=2)
        draw.ellipse((cx-265,y-33,cx+265,y+33),fill='white',outline='black',width=2)
        draw.multiline_text((cx,y),'\n'.join(textwrap.wrap(label,35)),font=font,anchor='mm',align='center',fill='black',spacing=1)
im.save(root/'use_cases_corrected.png')
for i,p in enumerate(d.paragraphs):
    if p.text.startswith('Figure 4:'):
        for prev in reversed(d.paragraphs[max(0,i-5):i]):
            nodes=prev._p.xpath('.//a:blip')
            if nodes:
                for node in nodes:d.part.related_parts[node.get(qn('r:embed'))]._blob=(root/'use_cases_corrected.png').read_bytes()
                break
d.styles['Normal'].font.size=Pt(12)
d.styles['Normal'].paragraph_format.line_spacing=1.15
for name,size in [('Heading 1',16),('Heading 2',14),('Heading 3',12)]:
    d.styles[name].font.size=Pt(size)
for p in d.paragraphs:
    if p.style.name=='Normal':
        p.paragraph_format.line_spacing=1.15
        p.paragraph_format.space_after=Pt(4)
    elif p.style.name.startswith('Heading'):
        p.paragraph_format.space_before=Pt(9)
        p.paragraph_format.space_after=Pt(5)
        for r in p.runs:r.font.color.rgb=RGBColor(0,0,0)
    if 'HOSPITAL MANAGEMENT SYSTEM'==p.text.strip():
        p.style=d.styles['Title']
        for r in p.runs:r.font.color.rgb=RGBColor(0,0,0)
for t in d.tables:
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.line_spacing=1.0
                p.paragraph_format.space_before=Pt(2)
                p.paragraph_format.space_after=Pt(2)
                for r in p.runs:r.font.size=Pt(10.5)
    pr=t._tbl.tblPr
    borders=pr.find(qn('w:tblBorders'))
    if borders is not None:
        for b in borders:b.set(qn('w:color'),'D9D9D9')
d.core_properties.title='Hospital Management System Project Documentation'
for p in d.paragraphs[11:38]:
    if not p.text.strip() and not p._p.xpath('.//w:drawing'):
        p.paragraph_format.space_after=Pt(0)
        p.paragraph_format.line_spacing=1
        for r in p.runs:r.font.size=Pt(6)
        p.paragraph_format.line_spacing=Pt(8)
for p in d.paragraphs:
    if p.text.strip()=='Acknowledgement':p.paragraph_format.page_break_before=True
    for r in p.runs:
        if '\ufffd' in r.text:r.text=r.text.replace('\ufffd','-')
for shape in d.inline_shapes:
    if shape.height>Pt(285):
        ratio=Pt(285)/shape.height
        shape.width=int(shape.width*ratio)
        shape.height=int(shape.height*ratio)
d.save(root/'HMS_Project_Documentation.docx')
print('Saved revised edition')
